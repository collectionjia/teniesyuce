# -*- coding: utf-8 -*-
"""Generate Word docs for OAI/Azure/GCP onboarding guides."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parents[1] / "docs"


def set_run_font(run, size=11, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if color:
        run.font.color.rgb = color


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=18, bold=True)
    p.paragraph_format.space_after = Pt(12)


def add_h(doc, text, level=1):
    p = doc.add_paragraph()
    sizes = {1: 14, 2: 12, 3: 11}
    run = p.add_run(text)
    set_run_font(run, size=sizes.get(level, 11), bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)


def add_p(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=11, bold=bold)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.35
    return p


def add_note(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=10, color=RGBColor(0x55, 0x55, 0x55))
    p.paragraph_format.space_after = Pt(8)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=11)
        p.paragraph_format.space_after = Pt(2)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(item)
        set_run_font(run, size=11)
        p.paragraph_format.space_after = Pt(2)


def set_cell_text(cell, text, bold=False, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=9)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            set_cell_text(table.rows[r_idx + 1].cells[c_idx], val, size=9)
    doc.add_paragraph()


def build_internal():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    add_title(doc, "官方 API（OAI）/ Azure / GCP：开通流程、要求与账号使用方式")
    add_note(
        doc,
        "整理说明：结合官方开通方式，以及当前商务侧对接口径（账号申请 → 定商务条件 → 开号 → 交付账号密码）。面向内部售前 / 交付，不替代各云厂商最新官方文档。",
    )

    add_h(doc, "一、三种方式总览", 1)
    add_table(
        doc,
        ["维度", "官方 OAI", "Azure", "GCP"],
        [
            ["入口", "platform.openai.com", "portal.azure.com / ai.azure.com", "console.cloud.google.com"],
            ["计费形态", "预充值（Credits / Billing）", "账号订阅制账单", "项目结算账号账单"],
            ["典型交付物", "API Key + 组织/项目权限", "账号+密码 + Endpoint + Key", "账号+密码 + Endpoint/Key/ADC"],
            ["中国大陆", "官方不可直接服务；无大陆官方折扣", "可通过合规企业订阅开通", "视结算主体、区域与模型而定"],
            ["折扣来源", "仅 OpenAI 官方；9折以下多属违规", "商务谈折扣/付款/税", "商务谈折扣/付款/税"],
            ["开通难度", "需海外支付与身份能力", "定商务条件后开号交付", "同 Azure"],
        ],
    )

    add_h(doc, "二、官方 OAI（OpenAI API Platform）", 1)
    add_h(doc, "2.1 是什么", 2)
    add_bullets(
        doc,
        [
            "OpenAI 官方 API 平台（常称 OAI）。",
            "入口：https://platform.openai.com/home",
            "属于 OpenAI 官方产品，不走 Azure / GCP 转售通道。",
        ],
    )
    add_h(doc, "2.2 开通流程（官方侧）", 2)
    add_numbered(
        doc,
        [
            "注册 / 登录 OpenAI 账号（组织 Organization）。",
            "完成身份与支付能力校验（海外手机、海外卡等，以官网当时要求为准）。",
            "在 Billing 中充值 / 绑定付款方式（预付费额度为主）。",
            "创建 Project，生成 API Key。",
            "按模型开通用量、设置限额与告警。",
        ],
    )
    add_h(doc, "2.3 要求", 2)
    add_table(
        doc,
        ["项", "说明"],
        [
            ["主体", "个人或企业组织；中国大陆主体通常无法直接走官方通路"],
            ["支付", "需可被官方接受的支付方式；依赖充值"],
            ["折扣", "折扣来自 OpenAI 官方；中国大陆没有官方折扣渠道"],
            ["海外折扣", "海外可有少量官方折扣；市场报价低于约9折，大概率为违规/黑产资源"],
            ["合规", "禁止转售未授权 Key、共享违规账号、绕过区域与 ToS 限制"],
        ],
    )
    add_h(doc, "2.4 账号使用方式", 2)
    add_bullets(
        doc,
        [
            "使用主体：客户持有 Platform 上的 Org / Project。",
            "调用方式：用 API Key 直连 api.openai.com（或以官方文档为准）。",
            "计费：扣平台预充值余额 / 账单；用完需再充。",
            "权限：Key 可按 Project 隔离；建议生产与测试分 Key，并设月度硬限额。",
            "交付形态：一般交付可用额度的官方账号体系 + Key，需注意 ToS 是否允许代持/转售。",
        ],
    )
    add_h(doc, "2.5 商务侧注意", 2)
    add_bullets(
        doc,
        [
            "适合：必须用 OpenAI 原生端点 / 最新模型同步更快的客户。",
            "不适合：无法海外支付、强合规要求只能走企业云合同的客户。",
            "风控：低于官方合理折扣的「OAI 批发价」默认按高风险资源处理。",
        ],
    )

    add_h(doc, "三、Azure（Azure OpenAI / Microsoft Foundry）", 1)
    add_h(doc, "3.1 是什么", 2)
    add_bullets(
        doc,
        [
            "在 Microsoft Azure 订阅内开通 Azure OpenAI / AI Foundry 等能力。",
            "管理入口：https://portal.azure.com/",
            "开发与模型管理：https://ai.azure.com/",
            "计费挂在 Azure 账号/订阅上，不是 OpenAI Platform 预充值。",
        ],
    )
    add_h(doc, "3.2 开通流程", 2)
    add_p(doc, "A. 官方自助（客户自己有订阅时）", bold=True)
    add_numbered(
        doc,
        [
            "注册 / 登录 Azure，创建订阅并绑卡或企业协议。",
            "创建资源组 → 创建 Azure OpenAI / AI 资源（选区域）。",
            "在 Foundry / Azure OpenAI 中部署模型（Deployment）。",
            "获取 Endpoint 与 API Key（或 Entra ID 鉴权）。",
            "按部署名调用；设置配额、网络与监控。",
        ],
    )
    add_p(doc, "B. 当前商务交付口径（代开 / 渠道开号）", bold=True)
    add_numbered(
        doc,
        [
            "客户侧向我方申请账号，并定好商务条件（折扣、付款、税等）。",
            "我方按条件开号。",
            "将账号 + 密码交付给客户（客户再用该账号进 Portal / Foundry 使用）。",
        ],
    )
    add_p(doc, "要点：Azure 是账号支付——费用走订阅账单；交付核心是「云账号」。")

    add_h(doc, "3.3 要求（商务条件）", 2)
    add_table(
        doc,
        ["条件", "说明"],
        [
            ["折扣", "相对公有价的商务折扣"],
            ["付款", "预付 / 账期等；最低付款额与账号体量相关"],
            ["税", "开票主体、税率、跨境税务安排"],
        ],
    )
    add_p(doc, "账号体量档位（美金 USD，对接口径示例）：")
    add_table(
        doc,
        ["档位", "含义"],
        [
            ["15K", "较小用量账号"],
            ["30K", "中等"],
            ["50K", "较大"],
            ["1M", "超大 / 企业级"],
        ],
    )

    add_h(doc, "3.4 账号使用方式", 2)
    add_bullets(
        doc,
        [
            "登录：客户使用交付的 Azure 账号登录 Portal / Foundry。",
            "资源：在订阅下查看资源、部署、密钥、项目 Endpoint。",
            "调用：Endpoint + Key，或企业场景用 Microsoft Entra ID。",
            "计费：用量计入该 Azure 订阅。",
            "权限建议：交付后尽快改密、开启 MFA、按最小权限建子账号。",
        ],
    )

    add_h(doc, "四、GCP（Vertex AI / Gemini API）", 1)
    add_h(doc, "4.1 是什么", 2)
    add_bullets(
        doc,
        [
            "在 Google Cloud 项目中开通 Vertex AI，调用 Gemini 等模型。",
            "入口：Google Cloud Console。",
            "计费挂在 GCP 结算账号 + 项目，与 Azure 同属云账号制。",
        ],
    )
    add_h(doc, "4.2 开通流程", 2)
    add_p(doc, "A. 官方自助", bold=True)
    add_numbered(
        doc,
        [
            "注册 Google 账号 / Cloud Identity，创建 GCP 项目。",
            "绑定 Billing Account。",
            "启用 Vertex AI API 等相关 API。",
            "配置区域、配额；创建服务账号或使用 API Key。",
            "通过 Vertex AI Endpoint / Gemini API 调用。",
        ],
    )
    add_p(doc, "B. 商务交付口径（与 Azure 同理）", bold=True)
    add_numbered(
        doc,
        [
            "申请账号并定商务条件（折扣、付款、税）。",
            "开号。",
            "交付账号 + 密码（及项目/结算相关信息）给客户。",
        ],
    )
    add_h(doc, "4.3 要求", 2)
    add_table(
        doc,
        ["项", "说明"],
        [
            ["结算", "有效 Billing Account；企业协议或预付视商务约定"],
            ["商务条件", "折扣 / 付款 / 税；额度档位可按客户交付规模谈"],
            ["区域与模型", "Gemini / 第三方模型可用性因 region 与项目资质而异"],
            ["合规", "遵守 Google Cloud / 模型提供方 AUP"],
        ],
    )
    add_h(doc, "4.4 账号使用方式", 2)
    add_bullets(
        doc,
        [
            "登录：Cloud Console 账号。",
            "调用：Vertex AI SDK / REST；鉴权常用服务账号 JSON。",
            "计费：项目用量 → 结算账号出账。",
            "权限建议：主账号仅作接管；日常用 IAM 最小权限服务账号。",
        ],
    )

    add_h(doc, "五、商务开号通用流程（Azure / GCP）", 1)
    add_p(
        doc,
        "客户申请账号 → 定商务条件（折扣/付款/税）→ 确认账号体量（如 15K/30K/50K/1M USD）→ 我方开号 → 交付账号密码 → 客户登录控制台部署模型/取 Endpoint 与 Key → 按约定充值或消耗额度。",
    )
    add_p(doc, "OAI 官方通道不走上述「云账号密码」主路径，而是 Platform 注册 + 充值 + API Key。")

    add_h(doc, "六、怎么选（简表）", 1)
    add_table(
        doc,
        ["客户诉求", "更合适"],
        [
            ["只要官方 OpenAI 原生 API、接受海外支付与预充值", "OAI"],
            ["要企业合同、发票、账号制、可谈折扣；接受 Azure 生态", "Azure"],
            ["要 Gemini / Google 生态或已有 GCP", "GCP"],
            ["只要很低的 OAI 折上折且来源不清", "不做（违规风险）"],
        ],
    )

    add_h(doc, "七、交付检查清单", 1)
    add_p(doc, "OAI", bold=True)
    add_bullets(
        doc,
        [
            "组织 / 项目已创建",
            "Billing 已充值且可扣费",
            "API Key 已生成并完成一次试调用",
            "限额与告警已设",
            "已向客户说明：非中国大陆官方折扣渠道；异常低价资源不可用",
        ],
    )
    add_p(doc, "Azure", bold=True)
    add_bullets(
        doc,
        [
            "订阅与账号密码已交付，客户可登录 Portal / Foundry",
            "订阅状态正常，额度 / 付款约定清晰",
            "模型已部署（或客户已知如何部署）",
            "Endpoint + Key（或 Entra 方案）已验证",
            "已提醒改密、MFA、权限拆分",
        ],
    )
    add_p(doc, "GCP", bold=True)
    add_bullets(
        doc,
        [
            "账号可登录 Console，项目与 Billing 已关联",
            "Vertex AI（及所需 API）已启用",
            "调用凭证已交付且试调用成功",
            "配额与区域符合客户模型需求",
            "已提醒改密、MFA、密钥轮换",
        ],
    )

    add_h(doc, "八、相关链接", 1)
    add_table(
        doc,
        ["平台", "链接"],
        [
            ["OpenAI Platform", "https://platform.openai.com/home"],
            ["Azure Portal", "https://portal.azure.com/"],
            ["Azure AI Foundry", "https://ai.azure.com/"],
            ["Google Cloud Console", "https://console.cloud.google.com/"],
        ],
    )

    add_h(doc, "修订记录", 1)
    add_table(doc, ["日期", "说明"], [["2026-03-18", "初版：汇总 OAI / Azure / GCP 开通、要求、使用方式及商务开号口径"]])

    path = OUT_DIR / "OAI-Azure-GCP-开通与账号使用说明.docx"
    doc.save(path)
    return path


def build_customer():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    add_title(doc, "OpenAI 官方 API / Azure / GCP 开通与使用说明（客户版）")
    add_note(
        doc,
        "本文面向客户，介绍三种主流大模型 API 开通方式的差异、开通要求与日常使用方式，便于选型与对接。具体商务条件（折扣、付款、税点、额度档位）以双方合同约定为准。",
    )

    add_h(doc, "一、三种方式对比", 1)
    add_table(
        doc,
        ["维度", "OpenAI 官方 API（OAI）", "Azure", "GCP"],
        [
            ["官方入口", "platform.openai.com", "portal.azure.com / ai.azure.com", "console.cloud.google.com"],
            ["计费方式", "平台预充值 / 账单扣费", "云订阅账号账单（按用量）", "云项目结算账号账单（按用量）"],
            ["您将获得", "组织/项目权限 + API Key", "云账号登录信息 + Endpoint + Key", "云账号登录信息 + 调用凭证"],
            ["适用场景", "需要直连 OpenAI 原生 API", "需要企业云合同、发票与账号管理", "需要 Gemini / Google 云生态"],
            ["中国大陆说明", "官方平台通常无法直接服务中国大陆主体", "可通过企业云订阅合规开通", "视结算主体与区域而定"],
        ],
    )

    add_h(doc, "二、OpenAI 官方 API（OAI）", 1)
    add_h(doc, "2.1 简介", 2)
    add_p(doc, "OpenAI 官方 API 平台，入口：https://platform.openai.com/home")
    add_p(doc, "调用走 OpenAI 原生接口，不经过 Azure / GCP。")
    add_h(doc, "2.2 开通流程", 2)
    add_numbered(
        doc,
        [
            "注册并登录 OpenAI 账号，创建组织（Organization）。",
            "按官网要求完成身份与支付方式验证。",
            "在 Billing 中充值或绑定付款方式。",
            "创建项目（Project），生成 API Key。",
            "配置用量限额与告警后即可调用。",
        ],
    )
    add_h(doc, "2.3 开通要求", 2)
    add_table(
        doc,
        ["项目", "说明"],
        [
            ["主体与支付", "需官网支持的主体与支付能力；中国大陆主体通常难以直接开通"],
            ["资金", "以预充值为主，余额不足将影响调用"],
            ["合规", "请仅使用官方正规渠道；异常低价的非官方资源存在封禁与合规风险，不建议使用"],
        ],
    )
    add_h(doc, "2.4 使用方式", 2)
    add_bullets(
        doc,
        [
            "使用 API Key 调用官方接口（Base URL 以官方文档为准）。",
            "费用从平台余额或账单中扣除。",
            "建议生产与测试环境使用不同 Key，并设置月度限额。",
        ],
    )

    add_h(doc, "三、Azure（Azure OpenAI / AI Foundry）", 1)
    add_h(doc, "3.1 简介", 2)
    add_bullets(
        doc,
        [
            "在 Microsoft Azure 订阅内使用 Azure OpenAI / AI Foundry。",
            "管理入口：https://portal.azure.com/",
            "模型与项目管理：https://ai.azure.com/",
            "费用记在 Azure 云账号订阅上，属于账号制使用。",
        ],
    )
    add_h(doc, "3.2 开通流程（由我方协助开通时）", 2)
    add_numbered(
        doc,
        [
            "您提出开通需求，并确认商务条件（折扣、付款方式、税务安排等）。",
            "确认账号额度规模（常见参考档位：15K / 30K / 50K / 1M 美金，具体以合同为准）。",
            "我方完成开号。",
            "向您交付账号与密码（以及必要的资源说明）。",
            "您登录 Azure Portal / AI Foundry，部署模型并获取 Endpoint 与 Key 后开始调用。",
        ],
    )
    add_p(doc, "若您已有自有 Azure 订阅，也可自行创建资源并部署模型；我方可提供技术协助。")
    add_h(doc, "3.3 开通要求", 2)
    add_table(
        doc,
        ["项目", "说明"],
        [
            ["商务条件", "需事先确认折扣、付款（预付/账期等）、开票与税务安排"],
            ["账号额度", "最低付款与可用规模取决于您需要的账号体量"],
            ["使用规范", "遵守 Microsoft 与模型提供方的服务条款与可接受使用政策"],
        ],
    )
    add_h(doc, "3.4 使用方式", 2)
    add_numbered(
        doc,
        [
            "使用交付的账号登录 Azure Portal 或 AI Foundry。",
            "在订阅中查看资源、模型部署、密钥与 Endpoint。",
            "通过 Endpoint + API Key 调用；企业场景也可配置 Microsoft Entra ID 鉴权。",
            "用量计入该订阅账单，按约定结算。",
        ],
    )
    add_p(doc, "安全建议：首次登录后请及时修改密码、开启 MFA，并为开发人员分配最小权限子账号。")

    add_h(doc, "四、GCP（Vertex AI / Gemini）", 1)
    add_h(doc, "4.1 简介", 2)
    add_bullets(
        doc,
        [
            "在 Google Cloud 项目中通过 Vertex AI 使用 Gemini 等模型。",
            "入口：https://console.cloud.google.com/",
            "计费挂在 GCP 结算账号与项目上，同属云账号制。",
        ],
    )
    add_h(doc, "4.2 开通流程（由我方协助开通时）", 2)
    add_numbered(
        doc,
        [
            "提出需求并确认商务条件（折扣、付款、税务）。",
            "确认账号/项目额度规模。",
            "我方开号并交付账号密码及相关项目信息。",
            "您登录 Google Cloud Console，启用 Vertex AI，配置调用凭证后开始使用。",
        ],
    )
    add_h(doc, "4.3 开通要求", 2)
    add_table(
        doc,
        ["项目", "说明"],
        [
            ["结算账号", "需有效 Billing Account，预付或企业协议以合同为准"],
            ["商务条件", "折扣、付款、税点需提前确认"],
            ["区域与模型", "不同区域可开通的模型与配额可能不同"],
        ],
    )
    add_h(doc, "4.4 使用方式", 2)
    add_bullets(
        doc,
        [
            "登录 Cloud Console 管理项目与配额。",
            "通过 Vertex AI SDK / REST 调用；常用服务账号密钥鉴权。",
            "费用按项目用量计入结算账号。",
            "建议主账号仅作管理用途，日常调用使用独立服务账号。",
        ],
    )

    add_h(doc, "五、协作开通流程（Azure / GCP）", 1)
    add_p(
        doc,
        "提出开通需求 → 确认商务条件（折扣/付款/税）→ 确认账号额度规模 → 完成开号 → 交付账号与密码 → 您登录控制台部署/启用模型并获取调用凭证 → 按约定充值或消耗额度。",
    )
    add_p(doc, "OpenAI 官方 API 一般不走「云账号密码」交付，而是平台账号 + 预充值 + API Key。")

    add_h(doc, "六、如何选择", 1)
    add_table(
        doc,
        ["您的需求", "建议方向"],
        [
            ["需要 OpenAI 原生 API，且具备海外支付与平台充值条件", "OpenAI 官方 API"],
            ["需要企业合同、发票、账号统一管理，接受 Azure 生态", "Azure"],
            ["需要 Gemini 或已有 Google Cloud 环境", "GCP"],
            ["来源不明、价格异常偏低的非官方渠道", "不建议，存在封号与合规风险"],
        ],
    )

    add_h(doc, "七、交付验收清单（供您核对）", 1)
    add_p(doc, "OpenAI 官方 API", bold=True)
    add_bullets(doc, ["组织 / 项目已就绪", "账单/余额可正常扣费", "API Key 已获取并完成试调用", "已设置用量限额与告警"])
    add_p(doc, "Azure", bold=True)
    add_bullets(
        doc,
        [
            "可使用账号密码登录 Portal / Foundry",
            "订阅状态正常，额度与付款约定清晰",
            "模型已部署（或已知部署步骤）",
            "Endpoint 与 Key（或企业鉴权）验证通过",
            "已完成改密与 MFA 等安全设置",
        ],
    )
    add_p(doc, "GCP", bold=True)
    add_bullets(
        doc,
        [
            "可登录 Console，项目已关联结算账号",
            "Vertex AI 等相关 API 已启用",
            "调用凭证已就绪并完成试调用",
            "区域与配额满足业务需求",
            "已完成安全与权限设置",
        ],
    )

    add_h(doc, "八、官方链接", 1)
    add_table(
        doc,
        ["平台", "链接"],
        [
            ["OpenAI Platform", "https://platform.openai.com/home"],
            ["Azure Portal", "https://portal.azure.com/"],
            ["Azure AI Foundry", "https://ai.azure.com/"],
            ["Google Cloud Console", "https://console.cloud.google.com/"],
        ],
    )
    add_p(doc, "如需报价、开号或技术对接，请联系您的商务对接人。具体条款以书面合同为准。")

    path = OUT_DIR / "OAI-Azure-GCP-开通与账号使用说明-客户版.docx"
    doc.save(path)
    return path


def main():
    try:
        from docx import Document  # noqa: F401
    except ImportError:
        import subprocess
        import sys

        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
    p1 = build_internal()
    p2 = build_customer()
    print(p1)
    print(p2)


if __name__ == "__main__":
    main()
